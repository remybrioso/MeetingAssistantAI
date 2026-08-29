"""
meeting_report_docx_exporter.py

Exporta MeetingReport como una minuta DOCX formal y legible.
"""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from models.artifacts.meeting_report import MeetingReport


class MeetingReportDocxExporter:
    """
    Convierte un MeetingReport validado en una minuta DOCX.

    El exportador es completamente determinista:
    no modifica el dominio, no reinterpreta contenido y no
    realiza llamadas a proveedores de IA.
    """

    FONT_NAME = "Arial"

    def export(
        self,
        report: MeetingReport,
        output_file: Path,
    ) -> None:
        """
        Construye y persiste la minuta DOCX.
        """

        self._validate_report(
            report
        )
        self._validate_output_file(
            output_file
        )

        document = self.build_document(
            report
        )

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document.save(
            output_file
        )

    def build_document(
        self,
        report: MeetingReport,
    ) -> DocumentType:
        """
        Construye el documento Word sin escribirlo en disco.

        Permite probar la representación de forma aislada.
        """

        self._validate_report(
            report
        )

        data = report.as_dict()

        document = Document()

        self._configure_document(
            document=document,
            data=data,
        )

        self._append_title(
            document=document,
            data=data,
        )
        self._append_objective(
            document=document,
            data=data,
        )
        self._append_executive_summary(
            document=document,
            data=data,
        )
        self._append_key_points(
            document=document,
            data=data,
        )
        self._append_topics(
            document=document,
            data=data,
        )
        self._append_decisions(
            document=document,
            data=data,
        )
        self._append_action_items(
            document=document,
            data=data,
        )
        self._append_risks(
            document=document,
            data=data,
        )
        self._append_pending_items(
            document=document,
            data=data,
        )
        self._append_participants(
            document=document,
            data=data,
        )
        self._append_conclusions(
            document=document,
            data=data,
        )
        self._append_footer(
            document
        )

        return document

    def _configure_document(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        section = document.sections[0]

        section.top_margin = Inches(
            0.7
        )
        section.bottom_margin = Inches(
            0.7
        )
        section.left_margin = Inches(
            0.8
        )
        section.right_margin = Inches(
            0.8
        )

        normal_style = document.styles[
            "Normal"
        ]
        normal_style.font.name = (
            self.FONT_NAME
        )
        normal_style.font.size = Pt(
            10.5
        )

        self._set_style_font(
            normal_style,
            self.FONT_NAME,
        )

        for style_name, size in (
            (
                "Title",
                22,
            ),
            (
                "Heading 1",
                15,
            ),
            (
                "Heading 2",
                12,
            ),
        ):
            style = document.styles[
                style_name
            ]
            style.font.name = (
                self.FONT_NAME
            )
            style.font.size = Pt(
                size
            )
            style.font.bold = True

            self._set_style_font(
                style,
                self.FONT_NAME,
            )

        core = (
            document.core_properties
        )
        core.title = data[
            "title"
        ]
        core.subject = (
            "Minuta de reunión generada por "
            "Meeting Assistant AI"
        )
        core.keywords = (
            "meeting, minutes, MAI"
        )

    def _append_title(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        label = document.add_paragraph()

        label.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        label.paragraph_format.space_after = Pt(
            4
        )

        run = label.add_run(
            "MINUTA DE REUNIÓN"
        )
        run.bold = True
        run.font.size = Pt(
            10
        )
        self._set_run_font(
            run,
            self.FONT_NAME,
        )

        title = document.add_paragraph(
            style="Title"
        )
        title.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )
        title.paragraph_format.space_after = Pt(
            6
        )

        title_run = title.add_run(
            data["title"]
        )
        self._set_run_font(
            title_run,
            self.FONT_NAME,
        )

        metadata = (
            document.add_paragraph()
        )
        metadata.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )
        metadata.paragraph_format.space_after = Pt(
            16
        )

        created_at = (
            self._format_created_at(
                data.get(
                    "created_at"
                )
            )
        )

        metadata_text = (
            "Generado por Meeting Assistant AI"
        )

        if created_at:
            metadata_text += (
                f" · {created_at}"
            )

        metadata_run = (
            metadata.add_run(
                metadata_text
            )
        )
        metadata_run.italic = True
        metadata_run.font.size = Pt(
            8.5
        )
        self._set_run_font(
            metadata_run,
            self.FONT_NAME,
        )

    def _append_objective(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Objetivo",
        )

        objective = data.get(
            "objective"
        )

        self._add_body_text(
            document,
            (
                objective
                if objective
                else (
                    "No se identificó un objetivo "
                    "explícito en la reunión."
                )
            ),
        )

    def _append_executive_summary(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Resumen ejecutivo",
        )

        self._add_body_text(
            document,
            data[
                "executive_summary"
            ],
        )

    def _append_key_points(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Puntos clave",
        )

        for point in data[
            "key_points"
        ]:
            paragraph = (
                document.add_paragraph(
                    style="List Bullet"
                )
            )
            paragraph.paragraph_format.space_after = Pt(
                3
            )

            run = paragraph.add_run(
                point
            )
            self._set_run_font(
                run,
                self.FONT_NAME,
            )

    def _append_topics(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Temas tratados",
        )

        topics = data.get(
            "topics",
            [],
        )

        if not topics:
            self._add_empty_message(
                document,
                "No se registraron temas adicionales.",
            )
            return

        for index, topic in enumerate(
            topics,
            start=1,
        ):
            self._add_subheading(
                document,
                (
                    f"{index}. "
                    f"{topic['title']}"
                ),
            )

            self._add_body_text(
                document,
                topic[
                    "summary"
                ],
            )

            self._append_evidence(
                document=document,
                evidence=topic.get(
                    "evidence",
                    [],
                ),
            )

    def _append_decisions(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Decisiones",
        )

        decisions = data.get(
            "decisions",
            [],
        )

        if not decisions:
            self._add_empty_message(
                document,
                "No se registraron decisiones.",
            )
            return

        for index, decision in enumerate(
            decisions,
            start=1,
        ):
            self._add_subheading(
                document,
                f"Decisión {index}",
            )

            self._add_body_text(
                document,
                decision[
                    "description"
                ],
            )

            rationale = decision.get(
                "rationale"
            )

            if rationale:
                self._add_labeled_text(
                    document=document,
                    label="Justificación",
                    value=rationale,
                )

            self._append_evidence(
                document=document,
                evidence=decision.get(
                    "evidence",
                    [],
                ),
            )

    def _append_action_items(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Acciones y compromisos",
        )

        action_items = data.get(
            "action_items",
            [],
        )

        if not action_items:
            self._add_empty_message(
                document,
                (
                    "No se registraron acciones "
                    "o compromisos."
                ),
            )
            return

        for index, action in enumerate(
            action_items,
            start=1,
        ):
            self._add_subheading(
                document,
                f"Acción {index}",
            )

            self._add_body_text(
                document,
                action[
                    "description"
                ],
            )

            owner = action.get(
                "owner"
            )

            owner_name = (
                owner.get(
                    "display_name"
                )
                if owner
                else None
            )

            metadata = [
                (
                    "Responsable",
                    (
                        owner_name
                        if owner_name
                        else "No asignado"
                    ),
                ),
                (
                    "Fecha límite",
                    (
                        action.get(
                            "due_date"
                        )
                        or "No definida"
                    ),
                ),
                (
                    "Estado",
                    self._format_status(
                        action.get(
                            "status"
                        )
                    ),
                ),
            ]

            self._add_metadata_table(
                document=document,
                values=metadata,
            )

            self._append_evidence(
                document=document,
                evidence=action.get(
                    "evidence",
                    [],
                ),
            )

    def _append_risks(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Riesgos y bloqueos",
        )

        risks = data.get(
            "risks",
            [],
        )

        if not risks:
            self._add_empty_message(
                document,
                (
                    "No se registraron riesgos "
                    "o bloqueos."
                ),
            )
            return

        for index, risk in enumerate(
            risks,
            start=1,
        ):
            self._add_subheading(
                document,
                f"Riesgo {index}",
            )

            self._add_body_text(
                document,
                risk[
                    "description"
                ],
            )

            impact = risk.get(
                "impact"
            )

            if impact:
                self._add_labeled_text(
                    document=document,
                    label="Impacto",
                    value=impact,
                )

            self._append_evidence(
                document=document,
                evidence=risk.get(
                    "evidence",
                    [],
                ),
            )

    def _append_pending_items(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Asuntos pendientes",
        )

        pending_items = data.get(
            "pending_items",
            [],
        )

        if not pending_items:
            self._add_empty_message(
                document,
                (
                    "No se registraron asuntos "
                    "pendientes."
                ),
            )
            return

        for index, pending in enumerate(
            pending_items,
            start=1,
        ):
            self._add_subheading(
                document,
                f"Pendiente {index}",
            )

            self._add_body_text(
                document,
                pending[
                    "description"
                ],
            )

            self._append_evidence(
                document=document,
                evidence=pending.get(
                    "evidence",
                    [],
                ),
            )

    def _append_participants(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Participantes",
        )

        participants = data.get(
            "participants",
            [],
        )

        if not participants:
            self._add_empty_message(
                document,
                (
                    "No fue posible identificar "
                    "participantes."
                ),
            )
            return

        for participant in participants:
            name = (
                participant.get(
                    "name"
                )
                or participant.get(
                    "speaker"
                )
                or "Participante"
            )

            details = []

            speaker = participant.get(
                "speaker"
            )
            role = participant.get(
                "role"
            )

            if speaker:
                details.append(
                    f"speaker: {speaker}"
                )

            if role:
                details.append(
                    f"rol: {role}"
                )

            text = name

            if details:
                text += (
                    " — "
                    + ", ".join(
                        details
                    )
                )

            paragraph = (
                document.add_paragraph(
                    style="List Bullet"
                )
            )

            run = paragraph.add_run(
                text
            )
            self._set_run_font(
                run,
                self.FONT_NAME,
            )

    def _append_conclusions(
        self,
        document: DocumentType,
        data: dict,
    ) -> None:
        self._add_heading(
            document,
            "Conclusiones",
        )

        conclusions = data.get(
            "conclusions",
            [],
        )

        if not conclusions:
            self._add_empty_message(
                document,
                (
                    "No se registraron conclusiones "
                    "explícitas."
                ),
            )
            return

        for conclusion in conclusions:
            paragraph = (
                document.add_paragraph(
                    style="List Bullet"
                )
            )

            run = paragraph.add_run(
                conclusion
            )
            self._set_run_font(
                run,
                self.FONT_NAME,
            )

    def _append_footer(
        self,
        document: DocumentType,
    ) -> None:
        section = document.sections[0]
        footer = section.footer

        paragraph = footer.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            (
                "Meeting Assistant AI · "
                "Documento generado automáticamente"
            )
        )
        run.font.size = Pt(
            8
        )
        self._set_run_font(
            run,
            self.FONT_NAME,
        )

    def _append_evidence(
        self,
        document: DocumentType,
        evidence: list[dict],
    ) -> None:
        if not evidence:
            return

        label = document.add_paragraph()
        label.paragraph_format.space_before = Pt(
            4
        )
        label.paragraph_format.space_after = Pt(
            2
        )

        run = label.add_run(
            "Evidencia"
        )
        run.bold = True
        run.font.size = Pt(
            9
        )
        self._set_run_font(
            run,
            self.FONT_NAME,
        )

        for reference in evidence:
            start = self._format_timestamp(
                reference[
                    "start"
                ]
            )
            end = self._format_timestamp(
                reference[
                    "end"
                ]
            )

            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(
                0.2
            )
            paragraph.paragraph_format.space_after = Pt(
                4
            )

            prefix = paragraph.add_run(
                (
                    f"{start}–{end} · "
                    f"{reference['speaker']}: "
                )
            )
            prefix.bold = True
            prefix.font.size = Pt(
                8.5
            )
            self._set_run_font(
                prefix,
                self.FONT_NAME,
            )

            excerpt = paragraph.add_run(
                f"“{reference['excerpt']}”"
            )
            excerpt.italic = True
            excerpt.font.size = Pt(
                8.5
            )
            self._set_run_font(
                excerpt,
                self.FONT_NAME,
            )

            self._shade_paragraph(
                paragraph,
                fill="F2F2F2",
            )

    def _add_heading(
        self,
        document: DocumentType,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph(
            style="Heading 1"
        )
        paragraph.paragraph_format.space_before = Pt(
            12
        )
        paragraph.paragraph_format.space_after = Pt(
            5
        )
        paragraph.paragraph_format.keep_with_next = True

        run = paragraph.add_run(
            text
        )
        self._set_run_font(
            run,
            self.FONT_NAME,
        )

    def _add_subheading(
        self,
        document: DocumentType,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph(
            style="Heading 2"
        )
        paragraph.paragraph_format.space_before = Pt(
            8
        )
        paragraph.paragraph_format.space_after = Pt(
            3
        )
        paragraph.paragraph_format.keep_with_next = True

        run = paragraph.add_run(
            text
        )
        self._set_run_font(
            run,
            self.FONT_NAME,
        )

    def _add_body_text(
        self,
        document: DocumentType,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(
            6
        )
        paragraph.paragraph_format.line_spacing = 1.08

        run = paragraph.add_run(
            text
        )
        self._set_run_font(
            run,
            self.FONT_NAME,
        )

    def _add_empty_message(
        self,
        document: DocumentType,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(
            6
        )

        run = paragraph.add_run(
            text
        )
        run.italic = True
        self._set_run_font(
            run,
            self.FONT_NAME,
        )

    def _add_labeled_text(
        self,
        document: DocumentType,
        label: str,
        value: str,
    ) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(
            5
        )

        label_run = paragraph.add_run(
            f"{label}: "
        )
        label_run.bold = True
        self._set_run_font(
            label_run,
            self.FONT_NAME,
        )

        value_run = paragraph.add_run(
            value
        )
        self._set_run_font(
            value_run,
            self.FONT_NAME,
        )

    def _add_metadata_table(
        self,
        document: DocumentType,
        values: list[
            tuple[str, str]
        ],
    ) -> None:
        table = document.add_table(
            rows=1,
            cols=len(
                values
            ),
        )

        table.autofit = True

        try:
            table.style = (
                "Light Shading Accent 1"
            )
        except KeyError:
            table.style = "Table Grid"

        for index, (
            label,
            value,
        ) in enumerate(
            values
        ):
            cell = table.rows[0].cells[
                index
            ]
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
            )

            label_run = paragraph.add_run(
                f"{label}\n"
            )
            label_run.bold = True
            label_run.font.size = Pt(
                8
            )
            self._set_run_font(
                label_run,
                self.FONT_NAME,
            )

            value_run = paragraph.add_run(
                value
            )
            value_run.font.size = Pt(
                9
            )
            self._set_run_font(
                value_run,
                self.FONT_NAME,
            )

    @staticmethod
    def _format_created_at(
        value: str | None,
    ) -> str | None:
        if not value:
            return None

        try:
            parsed = datetime.fromisoformat(
                value.replace(
                    "Z",
                    "+00:00",
                )
            )
        except ValueError:
            return value

        timezone_name = (
            parsed.tzname()
            if parsed.tzinfo
            else None
        )

        formatted = parsed.strftime(
            "%d/%m/%Y %H:%M"
        )

        if timezone_name:
            formatted += (
                f" {timezone_name}"
            )

        return formatted

    @staticmethod
    def _format_timestamp(
        seconds: float,
    ) -> str:
        total_seconds = int(
            seconds
        )

        hours, remainder = divmod(
            total_seconds,
            3600,
        )

        minutes, seconds_value = divmod(
            remainder,
            60,
        )

        return (
            f"{hours:02d}:"
            f"{minutes:02d}:"
            f"{seconds_value:02d}"
        )

    @staticmethod
    def _format_status(
        status: str | None,
    ) -> str:
        labels = {
            "unknown": "No especificado",
            "pending": "Pendiente",
            "completed": "Completado",
            "cancelled": "Cancelado",
        }

        return labels.get(
            status,
            "No especificado",
        )

    @staticmethod
    def _validate_report(
        report,
    ) -> None:
        if not isinstance(
            report,
            MeetingReport,
        ):
            raise TypeError(
                "report debe ser una instancia "
                "de MeetingReport."
            )

    @staticmethod
    def _validate_output_file(
        output_file,
    ) -> None:
        if not isinstance(
            output_file,
            Path,
        ):
            raise TypeError(
                "output_file debe ser una instancia "
                "de Path."
            )

        if (
            output_file.suffix.lower()
            != ".docx"
        ):
            raise ValueError(
                "output_file debe usar extensión .docx."
            )

    @staticmethod
    def _set_style_font(
        style,
        font_name: str,
    ) -> None:
        style.element.rPr.rFonts.set(
            qn(
                "w:ascii"
            ),
            font_name,
        )
        style.element.rPr.rFonts.set(
            qn(
                "w:hAnsi"
            ),
            font_name,
        )

    @staticmethod
    def _set_run_font(
        run,
        font_name: str,
    ) -> None:
        run.font.name = font_name

        run_properties = (
            run._element.get_or_add_rPr()
        )
        run_fonts = (
            run_properties.rFonts
        )

        if run_fonts is None:
            run_fonts = OxmlElement(
                "w:rFonts"
            )
            run_properties.insert(
                0,
                run_fonts,
            )

        run_fonts.set(
            qn(
                "w:ascii"
            ),
            font_name,
        )
        run_fonts.set(
            qn(
                "w:hAnsi"
            ),
            font_name,
        )

    @staticmethod
    def _shade_paragraph(
        paragraph,
        fill: str,
    ) -> None:
        paragraph_properties = (
            paragraph._p.get_or_add_pPr()
        )

        shading = paragraph_properties.find(
            qn(
                "w:shd"
            )
        )

        if shading is None:
            shading = OxmlElement(
                "w:shd"
            )
            paragraph_properties.append(
                shading
            )

        shading.set(
            qn(
                "w:fill"
            ),
            fill,
        )
