import json
from pathlib import Path

import pytest
from docx import Document

from application.dependency_container import (
    meeting_pipeline,
)
from models.artifacts.meeting_report import MeetingReport
from models.recording_session import RecordingSession
from models.transcript import Segment, Transcript
from services.transcript_storage_service import (
    TranscriptStorageService,
)
from services.workspace_service import WorkspaceService


def _build_e2e_transcript() -> Transcript:
    """
    Construye una reunión sintética autosuficiente con los cinco
    tipos semánticos principales utilizados por MeetingReport.
    """

    transcript = Transcript()

    segments = [
        Segment(
            start=0.0,
            end=8.0,
            speaker="LOCAL",
            text=(
                "El equipo decidió migrar la base de datos "
                "principal a la nube para reducir la dependencia "
                "de los servidores locales."
            ),
        ),
        Segment(
            start=8.0,
            end=16.0,
            speaker="REMOTE",
            text=(
                "María preparará el plan detallado de migración "
                "y deberá entregarlo a más tardar el "
                "30 de agosto de 2026."
            ),
        ),
        Segment(
            start=16.0,
            end=24.0,
            speaker="LOCAL",
            text=(
                "Existe el riesgo de una interrupción temporal "
                "del servicio durante la ejecución de la "
                "migración de la base de datos."
            ),
        ),
        Segment(
            start=24.0,
            end=32.0,
            speaker="REMOTE",
            text=(
                "Queda pendiente confirmar con Operaciones la "
                "ventana definitiva de mantenimiento antes de "
                "ejecutar el cambio."
            ),
        ),
        Segment(
            start=32.0,
            end=40.0,
            speaker="LOCAL",
            text=(
                "También revisamos la arquitectura actual, la "
                "capacidad disponible de los servidores y las "
                "dependencias técnicas existentes."
            ),
        ),
    ]

    for segment in segments:
        transcript.add_segment(
            segment
        )

    return transcript


def _build_recording_session(
    tmp_path: Path,
) -> RecordingSession:
    session = RecordingSession(
        base_output_dir=str(
            tmp_path
        ),
        session_prefix="meeting_e2e",
    )

    workspace = WorkspaceService().create(
        session.session_dir
    )

    session.attach_workspace(
        workspace
    )

    TranscriptStorageService().save(
        _build_e2e_transcript(),
        workspace.transcript_json,
    )

    return session


def _docx_text(
    path: Path,
) -> str:
    document = Document(
        path
    )

    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(
                paragraph.text
            )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(
                        cell.text
                    )

    return "\n".join(
        parts
    )


@pytest.mark.integration
def test_meeting_pipeline_generates_complete_artifact_delivery(
    tmp_path: Path,
) -> None:
    """
    E2E productivo autosuficiente.

    Ejecuta el MeetingPipeline real registrado en el contenedor,
    incluyendo extracción staged, consolidación global staged y
    entrega física de todos los artefactos derivados.

    Esta prueba requiere que Ollama y qwen2.5:3b estén disponibles.
    """

    session = _build_recording_session(
        tmp_path
    )

    report = meeting_pipeline.process(
        session
    )

    assert isinstance(
        report,
        MeetingReport,
    )

    assert (
        report.artifact_type
        == "meeting_report"
    )

    assert (
        report.prompt_version
        == "meeting_report_global_staged_v1"
    )

    workspace = session.workspace

    required_paths = [
        workspace.meeting_report_json,
        workspace.action_items_json,
        workspace.decisions_json,
        workspace.meeting_report_markdown,
        workspace.meeting_minutes_docx,
        workspace.meeting_pdf,
    ]

    assert all(
        path.exists()
        for path in required_paths
    )

    assert all(
        path.stat().st_size > 0
        for path in required_paths
    )

    report_json = json.loads(
        workspace.meeting_report_json.read_text(
            encoding="utf-8"
        )
    )

    action_items_json = json.loads(
        workspace.action_items_json.read_text(
            encoding="utf-8"
        )
    )

    decisions_json = json.loads(
        workspace.decisions_json.read_text(
            encoding="utf-8"
        )
    )

    assert (
        report_json
        == report.as_dict()
    )

    assert (
        action_items_json["artifact_type"]
        == "meeting_action_items"
    )

    assert (
        action_items_json["meeting_title"]
        == report.title
    )

    assert (
        action_items_json["provider"]
        == report.provider
    )

    assert (
        action_items_json["model"]
        == report.model
    )

    assert (
        action_items_json["prompt_version"]
        == report.prompt_version
    )

    assert (
        action_items_json["items"]
        == [
            item.as_dict()
            for item in report.action_items
        ]
    )

    assert (
        decisions_json["artifact_type"]
        == "meeting_decisions"
    )

    assert (
        decisions_json["meeting_title"]
        == report.title
    )

    assert (
        decisions_json["provider"]
        == report.provider
    )

    assert (
        decisions_json["model"]
        == report.model
    )

    assert (
        decisions_json["prompt_version"]
        == report.prompt_version
    )

    assert (
        decisions_json["items"]
        == [
            item.as_dict()
            for item in report.decisions
        ]
    )

    markdown = (
        workspace
        .meeting_report_markdown
        .read_text(
            encoding="utf-8"
        )
    )

    assert (
        f"# {report.title}"
        in markdown
    )

    assert (
        "## Resumen ejecutivo"
        in markdown
    )

    docx_text = _docx_text(
        workspace.meeting_minutes_docx
    )

    assert report.title in docx_text

    assert (
        "Resumen ejecutivo"
        in docx_text
    )

    assert (
        "Acciones y compromisos"
        in docx_text
    )

    pdf_prefix = (
        workspace
        .meeting_pdf
        .read_bytes()[:5]
    )

    assert pdf_prefix == b"%PDF-"

    assert (
        report.action_items
    )

    assert (
        report.decisions
    )

    assert all(
        item.evidence
        for item in report.action_items
    )

    assert all(
        item.evidence
        for item in report.decisions
    )
