from datetime import UTC, date, datetime
from pathlib import Path

import pytest
from docx import Document

from models.artifacts.meeting_report import MeetingReport
from models.meeting_report import (
    ActionItem,
    ActionStatus,
    Decision,
    EvidenceReference,
    MeetingRisk,
    MeetingTopic,
    Participant,
    PendingItem,
)
from services.meeting_report_docx_exporter import (
    MeetingReportDocxExporter,
)


def build_report() -> MeetingReport:
    evidence = EvidenceReference(
        speaker="SPEAKER_01",
        start=65.0,
        end=72.0,
        excerpt=(
            "El equipo DBA completará las pruebas "
            "antes del viernes."
        ),
    )

    return MeetingReport(
        artifact_type="meeting_report",
        provider="ollama",
        model="qwen2.5:3b",
        prompt_version=(
            "meeting_report_global_staged_v1"
        ),
        created_at=datetime(
            2026,
            8,
            29,
            0,
            40,
            tzinfo=UTC,
        ),
        title="Seguimiento del proyecto Oracle",
        objective=(
            "Revisar el avance y definir próximos pasos."
        ),
        executive_summary=(
            "Se revisó el avance del proyecto Oracle, "
            "los riesgos actuales y las acciones necesarias "
            "antes del próximo despliegue."
        ),
        key_points=[
            "Las pruebas de integración continúan.",
            "Se confirmó un posible retraso.",
            "El equipo DBA completará las pruebas.",
        ],
        topics=[
            MeetingTopic(
                title="Pruebas de integración",
                summary=(
                    "Se revisó el estado de las pruebas "
                    "previas al despliegue."
                ),
                evidence=[
                    evidence,
                ],
            ),
        ],
        decisions=[
            Decision(
                description=(
                    "Mantener el despliegue condicionado "
                    "a completar las pruebas."
                ),
                rationale=(
                    "Las pruebas todavía no han finalizado."
                ),
                evidence=[
                    evidence,
                ],
            ),
        ],
        action_items=[
            ActionItem(
                description=(
                    "Completar las pruebas de integración."
                ),
                owner="Equipo DBA",
                due_date=date(
                    2026,
                    8,
                    14,
                ),
                status=ActionStatus.PENDING,
                evidence=[
                    evidence,
                ],
            ),
        ],
        risks=[
            MeetingRisk(
                description=(
                    "Las pruebas podrían retrasar "
                    "el despliegue."
                ),
                impact=(
                    "La fecha de implementación podría "
                    "modificarse."
                ),
                evidence=[
                    evidence,
                ],
            ),
        ],
        pending_items=[
            PendingItem(
                description=(
                    "Confirmar la ventana definitiva "
                    "de mantenimiento."
                ),
                evidence=[
                    evidence,
                ],
            ),
        ],
        participants=[
            Participant(
                name="Remy",
                speaker="SPEAKER_00",
                role="Coordinador",
            ),
        ],
        conclusions=[
            (
                "El despliegue dependerá del resultado "
                "de las pruebas."
            ),
        ],
    )


def document_text(document) -> str:
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(
                paragraph.text
            )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(
                        cell.text
                    )

    return "\n".join(
        parts
    )


def test_docx_exporter_builds_complete_document() -> None:
    document = (
        MeetingReportDocxExporter()
        .build_document(
            build_report()
        )
    )

    content = document_text(
        document
    )

    expected_sections = [
        "MINUTA DE REUNIÓN",
        "Seguimiento del proyecto Oracle",
        "Objetivo",
        "Resumen ejecutivo",
        "Puntos clave",
        "Temas tratados",
        "Decisiones",
        "Acciones y compromisos",
        "Riesgos y bloqueos",
        "Asuntos pendientes",
        "Participantes",
        "Conclusiones",
    ]

    for expected in expected_sections:
        assert expected in content


def test_docx_exporter_renders_action_metadata() -> None:
    document = (
        MeetingReportDocxExporter()
        .build_document(
            build_report()
        )
    )

    content = document_text(
        document
    )

    assert (
        "Responsable\nEquipo DBA"
        in content
    )
    assert (
        "Fecha límite\n2026-08-14"
        in content
    )
    assert (
        "Estado\nPendiente"
        in content
    )


def test_docx_exporter_renders_evidence() -> None:
    document = (
        MeetingReportDocxExporter()
        .build_document(
            build_report()
        )
    )

    content = document_text(
        document
    )

    assert "Evidencia" in content
    assert (
        "00:01:05–00:01:12"
        in content
    )
    assert "SPEAKER_01" in content
    assert (
        "El equipo DBA completará las pruebas"
        in content
    )


def test_docx_exporter_renders_participants() -> None:
    document = (
        MeetingReportDocxExporter()
        .build_document(
            build_report()
        )
    )

    content = document_text(
        document
    )

    assert "Remy" in content
    assert "SPEAKER_00" in content
    assert "Coordinador" in content


def test_docx_exporter_writes_docx_file(
    tmp_path: Path,
) -> None:
    output_file = (
        tmp_path
        / "documents"
        / "minutes.docx"
    )

    MeetingReportDocxExporter().export(
        build_report(),
        output_file,
    )

    assert output_file.exists()
    assert output_file.stat().st_size > 0

    document = Document(
        output_file
    )

    assert (
        "Seguimiento del proyecto Oracle"
        in document_text(
            document
        )
    )


def test_docx_exporter_creates_parent_directory(
    tmp_path: Path,
) -> None:
    output_file = (
        tmp_path
        / "nested"
        / "documents"
        / "minutes.docx"
    )

    MeetingReportDocxExporter().export(
        build_report(),
        output_file,
    )

    assert output_file.exists()


def test_docx_exporter_sets_document_metadata() -> None:
    document = (
        MeetingReportDocxExporter()
        .build_document(
            build_report()
        )
    )

    assert (
        document.core_properties.title
        == "Seguimiento del proyecto Oracle"
    )
    assert (
        "Meeting Assistant AI"
        in document.core_properties.subject
    )


def test_docx_exporter_handles_empty_optional_sections() -> None:
    report = MeetingReport(
        artifact_type="meeting_report",
        provider="ollama",
        model="qwen2.5:3b",
        prompt_version=(
            "meeting_report_global_staged_v1"
        ),
        title="Reunión técnica",
        executive_summary=(
            "Se revisó el estado general del proyecto "
            "durante la reunión técnica."
        ),
        key_points=[
            "Se revisó el estado general del proyecto.",
        ],
    )

    document = (
        MeetingReportDocxExporter()
        .build_document(
            report
        )
    )

    content = document_text(
        document
    )

    expected_messages = [
        (
            "No se identificó un objetivo "
            "explícito en la reunión."
        ),
        "No se registraron temas adicionales.",
        "No se registraron decisiones.",
        (
            "No se registraron acciones "
            "o compromisos."
        ),
        (
            "No se registraron riesgos "
            "o bloqueos."
        ),
        (
            "No se registraron asuntos "
            "pendientes."
        ),
        (
            "No fue posible identificar "
            "participantes."
        ),
        (
            "No se registraron conclusiones "
            "explícitas."
        ),
    ]

    for expected in expected_messages:
        assert expected in content


def test_docx_exporter_does_not_mutate_report() -> None:
    report = build_report()
    original = report.as_dict()

    (
        MeetingReportDocxExporter()
        .build_document(
            report
        )
    )

    assert report.as_dict() == original


def test_docx_exporter_rejects_invalid_report() -> None:
    with pytest.raises(
        TypeError,
        match="MeetingReport",
    ):
        (
            MeetingReportDocxExporter()
            .build_document(
                object()
            )
        )


def test_docx_exporter_rejects_invalid_output_type() -> None:
    with pytest.raises(
        TypeError,
        match="Path",
    ):
        MeetingReportDocxExporter().export(
            build_report(),
            "minutes.docx",
        )


def test_docx_exporter_rejects_non_docx_extension(
    tmp_path: Path,
) -> None:
    with pytest.raises(
        ValueError,
        match=".docx",
    ):
        MeetingReportDocxExporter().export(
            build_report(),
            tmp_path / "minutes.pdf",
        )


@pytest.mark.parametrize(
    (
        "status",
        "expected",
    ),
    [
        (
            "unknown",
            "No especificado",
        ),
        (
            "pending",
            "Pendiente",
        ),
        (
            "completed",
            "Completado",
        ),
        (
            "cancelled",
            "Cancelado",
        ),
    ],
)
def test_docx_exporter_formats_action_status(
    status: str,
    expected: str,
) -> None:
    assert (
        MeetingReportDocxExporter
        ._format_status(
            status
        )
        == expected
    )


def test_docx_exporter_formats_unknown_status() -> None:
    assert (
        MeetingReportDocxExporter
        ._format_status(
            "other"
        )
        == "No especificado"
    )


def test_docx_exporter_formats_long_timestamp() -> None:
    assert (
        MeetingReportDocxExporter
        ._format_timestamp(
            3723.0
        )
        == "01:02:03"
    )



def test_docx_exporter_formats_created_at() -> None:
    assert (
        MeetingReportDocxExporter
        ._format_created_at(
            "2026-08-29T00:40:25+00:00"
        )
        == "29/08/2026 00:40 UTC"
    )


def test_docx_exporter_preserves_unparseable_created_at() -> None:
    assert (
        MeetingReportDocxExporter
        ._format_created_at(
            "fecha-desconocida"
        )
        == "fecha-desconocida"
    )
